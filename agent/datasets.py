"""Bounded, trusted dataset operations. Never executes uploaded formulas or code."""
import asyncio
import csv
import hashlib
import io
import json
import os
import sys
import uuid
import zipfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, getcontext
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLS = 100
MAX_CELLS = 200000
MAX_GROUPS = 500


class DataError(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def directory(root, dataset_id):
    # IDs, not uploaded filenames, determine every storage path.
    return Path(root) / "datasets" / str(uuid.UUID(str(dataset_id)))


def read_dataset(root, dataset_id):
    path = directory(root, dataset_id) / "table.json"
    if not path.is_file():
        raise DataError("missing_dataset", "Dataset not found. Upload the file again.")
    data = json.loads(path.read_text())
    if hashlib.sha256(json.dumps(data["rows"], ensure_ascii=True).encode()).hexdigest() != data["sha256"]:
        raise DataError("integrity", "Stored dataset integrity check failed. Upload it again.")
    return data


def metadata(data):
    return {k: v for k, v in data.items() if k != "rows"} | {"preview": data["rows"][:8]}


def number(value):
    if value is None:
        return None
    try:
        d = Decimal(str(value).strip())
        if not d.is_finite() or abs(d) > Decimal("1e15") or d.as_tuple().exponent < -10:
            return None
        return d
    except (InvalidOperation, ValueError):
        return None


def normalize(rows):
    iterator = iter(rows)
    try:
        header = [str(v).strip() if v is not None else "" for v in next(iterator)]
    except StopIteration:
        raise DataError("empty_file", "File has no header row.") from None
    if not header or len(header) > MAX_COLS or any(not h or len(h) > 100 for h in header):
        raise DataError("bad_header", "Use 1-100 named columns with headers of 1-100 characters.")
    if len(set(header)) != len(header):
        raise DataError("bad_header", "Duplicate column names are ambiguous. Rename them before upload.")
    output = []
    for line, row in enumerate(iterator, 2):
        if line > MAX_ROWS + 1 or (line - 1) * len(header) > MAX_CELLS:
            raise DataError("size_limit", "Maximum 20,000 rows and 200,000 cells per file.")
        if len(row) != len(header):
            raise DataError("ragged_rows", f"Row {line} has a different number of fields than the header.")
        values = []
        for v in row:
            if isinstance(v, (datetime, date)):
                v = v.isoformat()
            if v is not None:
                v = str(v)
                if len(v) > 2000:
                    raise DataError("cell_limit", f"Row {line} contains a cell longer than 2,000 characters.")
                if not v.strip():
                    v = None
            values.append(v)
        output.append(values)
    if not output:
        raise DataError("empty_file", "File has a header but no data rows.")
    return header, output


def ingest(root, dataset_id, options):
    folder = directory(root, dataset_id)
    raw = (folder / "source").read_bytes()
    if len(raw) > MAX_BYTES:
        raise DataError("size_limit", "Maximum upload size is 5 MiB.")
    suffix = Path(options["filename"]).suffix.lower()
    selected = options.get("sheet") or None
    sheets = []
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeError:
            raise DataError("encoding", "Save the CSV as UTF-8 and upload again.") from None
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        header, rows = normalize(csv.reader(io.StringIO(text), dialect, strict=True))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 25 * 1024**2:
                raise DataError("size_limit", "Excel archive expands beyond 25 MiB or contains too many parts.")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = wb.sheetnames
            selected = selected or sheets[0]
            if selected not in sheets:
                raise DataError("bad_sheet", "Worksheet not found. Available: " + ", ".join(sheets))
            ws = wb[selected]
            if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > MAX_COLS:
                raise DataError("size_limit", "Worksheet dimensions exceed the row/column limits.")

            def cells():
                for row in ws.iter_rows():
                    if any(c.data_type in ("f", "e") for c in row):
                        raise DataError("formula_cells", "Excel formulas/errors are not evaluated. Upload a values-only copy.")
                    yield [c.value for c in row]
            header, rows = normalize(cells())
        finally:
            wb.close()
    elif suffix == ".xls":
        raise DataError("unsupported_format", "Save legacy .xls files as .xlsx or UTF-8 CSV first.")
    else:
        raise DataError("unsupported_format", "Upload a .csv or .xlsx file.")
    columns = []
    for idx, name in enumerate(header):
        values = [row[idx] for row in rows]
        present = [v for v in values if v is not None]
        numeric = sum(number(v) is not None for v in present)
        columns.append({"name": name, "missing": len(values) - len(present),
                        "numeric_values": numeric,
                        "type": "number" if present and numeric == len(present) else "text"})
    data = {"id": dataset_id, "filename": Path(options["filename"]).name,
            "sheet": selected, "sheets": sheets, "row_count": len(rows),
            "columns": columns, "rows": rows,
            "sha256": hashlib.sha256(json.dumps(rows, ensure_ascii=True).encode()).hexdigest()}
    temp = folder / "table.tmp"
    temp.write_text(json.dumps(data, ensure_ascii=True))
    temp.replace(folder / "table.json")
    return metadata(data)


def profile(data):
    duplicates = len(data["rows"]) - len({tuple(row) for row in data["rows"]})
    missing = sum(c["missing"] for c in data["columns"])
    result = {"kind": "profile", "dataset_id": data["id"], "source_sha256": data["sha256"],
              "row_count": data["row_count"], "columns": data["columns"],
              "duplicate_rows": duplicates, "missing_cells": missing,
              "verified": True, "warnings": [],
              "summary": f"{data['row_count']} rows, {len(data['columns'])} columns; "
                         f"{missing} blank cells and {duplicates} exact duplicate rows. No rows were removed."}
    if duplicates:
        result["warnings"].append("Exact duplicates remain included; confirm the intended row grain before reporting totals.")
    if missing:
        result["warnings"].append("Blank values are retained and reported, never silently replaced with zero.")
    return result


def aggregate(data, options, trend=False):
    # 20k bounded values with ten decimal places fit exactly within 50 digits.
    getcontext().prec = 50
    names = [c["name"] for c in data["columns"]]
    metric = options.get("metric")
    group = options.get("date_column") if trend else options.get("group_by")
    for column in (metric, group):
        if column not in names:
            raise DataError("bad_column", f"Column {column!r} was not found. Choose an exact uploaded column name.")
    mi, gi = names.index(metric), names.index(group)
    buckets = defaultdict(list)
    source_values = []
    missing = 0
    for line, row in enumerate(data["rows"], 2):
        key = row[gi]
        if trend:
            try:
                # Require unambiguous ISO dates; no guessed month/day order.
                text = str(key)
                parsed = datetime.fromisoformat(text) if "T" in text or " " in text else date.fromisoformat(text)
                key = parsed.strftime("%Y-%m")
            except ValueError:
                raise DataError("bad_date", f"Row {line} has an invalid date. Use YYYY-MM-DD without blanks.") from None
        value = number(row[mi])
        if row[mi] is None:
            missing += 1
        elif value is None:
            raise DataError("bad_numeric", f"Row {line} in {metric} is not a finite plain number. Remove currency symbols/commas and correct text.")
        buckets[key].append(value)
        if value is not None:
            source_values.append(value)
        if len(buckets) > MAX_GROUPS:
            raise DataError("group_limit", "More than 500 groups. Choose a less detailed grouping column.")
    if not source_values:
        raise DataError("no_values", "The selected measure contains no numeric values.")
    table = []
    for key, values in buckets.items():
        valid = [v for v in values if v is not None]
        total = sum(valid, Decimal(0)) if valid else None
        table.append({"group": key, "rows": len(values), "valid_values": len(valid),
                      "missing_values": len(values) - len(valid),
                      "sum": str(total) if total is not None else None,
                      "mean": str(total / len(valid)) if valid else None})
    table.sort(key=lambda r: str(r["group"]))
    total = sum(source_values, Decimal(0))
    # Separate aggregation path for reconciliation; counts include blanks.
    check_total = sum((Decimal(r["sum"]) for r in table if r["sum"] is not None), Decimal(0))
    verified = check_total == total and sum(r["rows"] for r in table) == data["row_count"]
    warnings = profile(data)["warnings"]
    if missing:
        warnings.append(f"{missing} blank {metric} values excluded from sum and mean; all-blank groups show null.")
    if any(r["group"] is None for r in table):
        warnings.append("Missing group labels appear as a separate null group.")
    if trend:
        previous = None
        for item in table:
            current = Decimal(item["sum"]) if item["sum"] is not None else None
            item["change_percent"] = (str((current - previous) / abs(previous) * 100)
                                      if previous not in (None, 0) and current is not None else None)
            previous = current
        warnings.append("Change is versus the previous observed month. Missing months are not imputed; zero baselines return null.")
    return {"kind": "trend" if trend else "aggregate", "dataset_id": data["id"],
            "source_sha256": data["sha256"], "metric": metric, "group_by": group,
            "row_count": data["row_count"], "total": str(total), "table": table,
            "verified": verified, "warnings": warnings,
            "checks": {"source_total": str(total), "group_total": str(check_total),
                       "source_rows": data["row_count"], "group_rows": sum(r["rows"] for r in table)},
            "summary": f"Sum of {metric}: {total} across {len(table)} {'months' if trend else 'groups'}. "
                       f"{len(source_values)} valid values; {missing} blanks excluded. Group totals and row counts reconciled."}


async def run_operation(root, dataset_id, operation, options=None, timeout=15):
    """Isolate parsers/analysis from the API event loop and kill them at deadline."""
    directory(root, dataset_id)
    process = await asyncio.create_subprocess_exec(
        sys.executable, "-m", "agent.datasets", str(Path(root).resolve()), dataset_id, operation,
        json.dumps(options or {}), stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
        env={k: v for k, v in os.environ.items() if k in ("PATH", "SYSTEMROOT", "WINDIR", "LANG")})
    try:
        async with asyncio.timeout(timeout):
            output = bytearray()
            while chunk := await process.stdout.read(65536):
                output.extend(chunk)
                if len(output) > 2 * 1024 * 1024:
                    raise DataError("output_limit", "Dataset result exceeds output limit.")
            await process.wait()
            if process.returncode:
                raise DataError("worker_error", "Dataset worker stopped. Try a smaller file.")
            return json.loads(output)
    except TimeoutError:
        return {"ok": False, "error": {"code": "timeout", "message": "Dataset operation exceeded its deadline.", "retryable": True}}
    except (DataError, ValueError) as exc:
        return {"ok": False, "error": {"code": getattr(exc, "code", "worker_error"),
                "message": str(exc)[:350], "retryable": False}}
    finally:
        if process.returncode is None:
            process.kill()
        await process.wait()


def worker():
    root, dataset_id, operation, options = sys.argv[1:]
    try:
        if sys.platform == "linux":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024**2, 512 * 1024**2))
            resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
        options = json.loads(options)
        if operation == "ingest":
            result = ingest(root, dataset_id, options)
        else:
            data = read_dataset(root, dataset_id)
            if operation == "metadata":
                result = metadata(data)
            elif operation == "dataset_profile":
                result = profile(data)
            elif operation in ("dataset_aggregate", "dataset_trend"):
                result = aggregate(data, options, operation == "dataset_trend")
            else:
                raise DataError("bad_input", "Unsupported dataset operation.")
        output = {"ok": True, "data": result}
    except Exception as exc:
        output = {"ok": False, "error": {"code": getattr(exc, "code", "invalid_file"),
                  "message": str(exc)[:350] if isinstance(exc, DataError) else "File could not be parsed. Check format and size.",
                  "retryable": False}}
    print(json.dumps(output, allow_nan=False))


if __name__ == "__main__":
    worker()
